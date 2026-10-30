import re
import fitz
from docx import Document as WordDocument
from llama_index.core import Document, Settings
from llama_index.core.schema import BaseNode
from llama_index.core.node_parser import SentenceSplitter
from dotenv import load_dotenv
from pathlib import Path
from typing import Any, List
from tqdm import tqdm
from ...setting import RAGSettings

load_dotenv()


class LocalDataIngestion:
    def __init__(self, setting: RAGSettings | None = None) -> None:
        self._setting = setting or RAGSettings()
        self._node_store = {}
        self._ingested_file = []

    def _filter_text(self, text):
        # Define the regex pattern.
        pattern = r'[a-zA-Z0-9 \u00C0-\u01B0\u1EA0-\u1EF9`~!@#$%^&*()_\-+=\[\]{}|\\;:\'",.<>/?]+'
        matches = re.findall(pattern, text)
        # Join all matched substrings into a single string
        filtered_text = " ".join(matches)
        # Normalize the text by removing extra whitespaces
        normalized_text = re.sub(r"\s+", " ", filtered_text.strip())

        return normalized_text

    def _extract_text(self, input_file: str) -> str:
        path = Path(input_file)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            with fitz.open(input_file) as pdf:
                return " ".join(page.get_text("text") for page in pdf)
        if suffix == ".docx":
            document = WordDocument(input_file)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            return "\n".join(paragraphs)
        if suffix in {".txt", ".csv", ".md", ".json"}:
            return path.read_text(encoding="utf-8", errors="replace")
        raise ValueError(
            f"Unsupported document format '{suffix or '<none>'}' for '{path.name}'. "
            "Supported formats are PDF, DOCX, TXT, CSV, MD, and JSON."
        )

    def store_nodes(
        self,
        input_files: list[str],
        embed_nodes: bool = True,
        embed_model: Any | None = None,
    ) -> List[BaseNode]:
        return_nodes = []
        self._ingested_file = []
        if len(input_files) == 0:
            return return_nodes
        splitter = SentenceSplitter.from_defaults(
            chunk_size=self._setting.ingestion.chunk_size,
            chunk_overlap=self._setting.ingestion.chunk_overlap,
            paragraph_separator=self._setting.ingestion.paragraph_sep,
            secondary_chunking_regex=self._setting.ingestion.chunking_regex,
        )
        if embed_nodes:
            Settings.embed_model = embed_model or Settings.embed_model
        for input_file in tqdm(input_files, desc="Ingesting data"):
            file_name = Path(input_file).name
            self._ingested_file.append(file_name)
            if file_name in self._node_store:
                return_nodes.extend(self._node_store[file_name])
            else:
                all_text = self._filter_text(self._extract_text(input_file))
                document = Document(
                    text=all_text.strip(),
                    metadata={
                        "file_name": file_name,
                    },
                )

                nodes = splitter([document], show_progress=True)
                if embed_nodes:
                    nodes = Settings.embed_model(nodes, show_progress=True)
                self._node_store[file_name] = nodes
                return_nodes.extend(nodes)
        return return_nodes

    def reset(self):
        self._node_store = {}
        self._ingested_file = []

    def check_nodes_exist(self):
        return len(self._node_store.values()) > 0

    def get_all_nodes(self):
        return_nodes = []
        for nodes in self._node_store.values():
            return_nodes.extend(nodes)
        return return_nodes

    def get_ingested_nodes(self):
        return_nodes = []
        for file in self._ingested_file:
            return_nodes.extend(self._node_store[file])
        return return_nodes
